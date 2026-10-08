"""Generate AegisAI hackathon pitch deck (PPTX + PDF).

Run from repo root or docs/hackathon:
  python docs/hackathon/build_pitch_deck.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "pitch-assets"
OUT_PPTX = HERE / "AegisAI-Pitch-Deck.pptx"
OUT_PDF = HERE / "AegisAI-Pitch-Deck.pdf"

# Widescreen 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Palette — charcoal + teal (security, not purple AI cliché)
BG = RGBColor(0x0B, 0x12, 0x20)
BG_CARD = RGBColor(0x12, 0x1A, 0x2B)
ACCENT = RGBColor(0x2D, 0xD4, 0xBF)  # teal
ACCENT_2 = RGBColor(0x38, 0xBD, 0xF8)  # sky
WARN = RGBColor(0xFB, 0xBF, 0x24)  # amber
DANGER = RGBColor(0xFB, 0x71, 0x85)  # rose
TEXT = RGBColor(0xF1, 0xF5, 0xF9)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# reportlab RGB tuples 0-1
R_BG = (0x0B / 255, 0x12 / 255, 0x20 / 255)
R_CARD = (0x12 / 255, 0x1A / 255, 0x2B / 255)
R_ACCENT = (0x2D / 255, 0xD4 / 255, 0xBF / 255)
R_ACCENT2 = (0x38 / 255, 0xBD / 255, 0xF8 / 255)
R_TEXT = (0xF1 / 255, 0xF5 / 255, 0xF9 / 255)
R_MUTED = (0x94 / 255, 0xA3 / 255, 0xB8 / 255)
R_WARN = (0xFB / 255, 0xBF / 255, 0x24 / 255)
R_DANGER = (0xFB / 255, 0x71 / 255, 0x85 / 255)


def _set_run(run, size=18, bold=False, color=TEXT, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill_slide(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _round_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def _textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def _multitext(slide, left, top, width, height, lines, *, size=16, color=TEXT, bold=False, spacing=6):
    """lines: list[str] or list[tuple[str, dict]]"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, opts = item
        else:
            text, opts = item, {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(opts.get("spacing", spacing))
        run = p.add_run()
        run.text = text
        _set_run(
            run,
            size=opts.get("size", size),
            bold=opts.get("bold", bold),
            color=opts.get("color", color),
            font=opts.get("font", "Calibri"),
        )
    return box


def _picture(slide, path: Path, left, top, width=None, height=None):
    if not path.exists():
        return None
    if width and height:
        return slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    if width:
        return slide.shapes.add_picture(str(path), left, top, width=width)
    return slide.shapes.add_picture(str(path), left, top, height=height)


def _accent_bar(slide):
    _rect(slide, Inches(0), Inches(0), Inches(0.12), SLIDE_H, ACCENT)


def _footer(slide, page: int, total: int = 8):
    _textbox(
        slide,
        Inches(0.5),
        Inches(7.1),
        Inches(8),
        Inches(0.3),
        "AegisAI  ·  ET AI Hackathon 2026 — Agentic Edition  ·  Problem 2",
        size=10,
        color=MUTED,
    )
    _textbox(
        slide,
        Inches(11.5),
        Inches(7.1),
        Inches(1.5),
        Inches(0.3),
        f"{page} / {total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def build_pptx() -> Path:
    """Clean rebuild with a polished title slide (no overlapping mess)."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    total = 8

    # 1 Title
    s = prs.slides.add_slide(blank)
    _fill_slide(s)
    _picture(s, ASSETS / "circuit.jpg", Inches(7.5), Inches(0), width=Inches(5.9), height=SLIDE_H)
    _rect(s, Inches(0), Inches(0), Inches(8.0), SLIDE_H, BG)
    # Gradient-ish edge
    _rect(s, Inches(7.7), Inches(0), Inches(0.35), SLIDE_H, BG)
    _picture(s, ASSETS / "brand-mark.jpg", Inches(0.65), Inches(0.55), width=Inches(1.5))
    _textbox(s, Inches(0.65), Inches(2.35), Inches(6.6), Inches(0.9), "AegisAI", size=56, bold=True, color=WHITE)
    _textbox(
        s,
        Inches(0.65),
        Inches(3.25),
        Inches(6.6),
        Inches(0.45),
        "Agentic Prompt Injection Firewall",
        size=22,
        bold=True,
        color=ACCENT,
    )
    _textbox(
        s,
        Inches(0.65),
        Inches(3.9),
        Inches(6.4),
        Inches(0.9),
        "A runtime security layer between untrusted content\nand agent tool actions.",
        size=16,
        color=MUTED,
    )
    _rect(s, Inches(0.65), Inches(5.1), Inches(1.4), Inches(0.07), ACCENT)
    _textbox(
        s,
        Inches(0.65),
        Inches(5.4),
        Inches(6.5),
        Inches(0.9),
        "ET AI Hackathon 2026 — Agentic Edition\nProblem 2  ·  Pitch Deck",
        size=14,
        color=MUTED,
    )
    _textbox(s, Inches(8.2), Inches(6.9), Inches(4.7), Inches(0.3), "Imagery: Unsplash  ·  Brand mark: generated", size=9, color=MUTED)

    # 2 Problem
    s = prs.slides.add_slide(blank)
    _fill_slide(s)
    _accent_bar(s)
    _round_rect(s, Inches(8.3), Inches(1.4), Inches(4.5), Inches(4.6), BG_CARD)
    _picture(s, ASSETS / "code.jpg", Inches(8.45), Inches(1.55), width=Inches(4.2), height=Inches(4.3))
    _textbox(s, Inches(0.5), Inches(0.35), Inches(7), Inches(0.35), "THE PROBLEM", size=12, bold=True, color=ACCENT)
    _textbox(s, Inches(0.5), Inches(0.75), Inches(7.5), Inches(1.0), "Agents read more than chat.", size=34, bold=True, color=WHITE)
    _multitext(
        s,
        Inches(0.5),
        Inches(1.9),
        Inches(7.3),
        Inches(4.6),
        [
            ("AI agents ingest documents, websites, emails, and tool outputs.", {"size": 17, "color": MUTED}),
            (" ", {"size": 8}),
            ("Hidden instructions can try to:", {"size": 17, "color": TEXT, "bold": True}),
            (" ", {"size": 6}),
            ("→  Override system intent (prompt injection)", {"size": 16, "color": TEXT}),
            ("→  Hijack tool calls (email, delete, shell)", {"size": 16, "color": TEXT}),
            ("→  Steal secrets or credentials", {"size": 16, "color": TEXT}),
            ("→  Poison MCP tool definitions", {"size": 16, "color": TEXT}),
            (" ", {"size": 10}),
            ("Detection alone is not enough — a compromised plan\nmust still be stopped before execution.", {"size": 15, "color": WARN, "bold": True}),
        ],
    )
    _footer(s, 2, total)

    # 3 Solution
    s = prs.slides.add_slide(blank)
    _fill_slide(s)
    _accent_bar(s)
    _textbox(s, Inches(0.5), Inches(0.35), Inches(10), Inches(0.3), "THE SOLUTION", size=12, bold=True, color=ACCENT)
    _textbox(s, Inches(0.5), Inches(0.7), Inches(12), Inches(0.6), "AegisAI sits between untrusted content and agent actions.", size=26, bold=True, color=WHITE)
    _textbox(
        s,
        Inches(0.5),
        Inches(1.4),
        Inches(12),
        Inches(0.4),
        "Runtime security layer — the frontend never decides ALLOW / BLOCK.",
        size=15,
        color=MUTED,
    )
    stages = [
        ("DETECT", "Rules + Prompt Guard\n+ Semantic Safeguard"),
        ("UNDERSTAND", "Evidence fusion\n& risk scoring"),
        ("DECIDE", "ALLOW · REVIEW\n· BLOCK policy"),
        ("PROTECT", "Authz + Tool Firewall\n+ MCP gateway"),
        ("AUDIT", "Hashed events\nno raw prompts"),
    ]
    x = 0.5
    for i, (title, body) in enumerate(stages):
        _round_rect(s, Inches(x), Inches(2.2), Inches(2.3), Inches(3.5), BG_CARD)
        _rect(s, Inches(x), Inches(2.2), Inches(2.3), Inches(0.1), ACCENT if i % 2 == 0 else ACCENT_2)
        _textbox(s, Inches(x + 0.15), Inches(2.5), Inches(2.0), Inches(0.4), f"0{i+1}", size=13, bold=True, color=ACCENT)
        _textbox(s, Inches(x + 0.15), Inches(3.05), Inches(2.0), Inches(0.45), title, size=15, bold=True, color=WHITE)
        _textbox(s, Inches(x + 0.15), Inches(3.7), Inches(2.0), Inches(1.3), body, size=13, color=MUTED)
        x += 2.5
    _footer(s, 3, total)

    # 4 Architecture
    s = prs.slides.add_slide(blank)
    _fill_slide(s)
    _accent_bar(s)
    _textbox(s, Inches(0.5), Inches(0.3), Inches(10), Inches(0.3), "ARCHITECTURE", size=12, bold=True, color=ACCENT)
    _textbox(s, Inches(0.5), Inches(0.6), Inches(12), Inches(0.5), "Defense in depth — detect, then stop the tool boundary.", size=24, bold=True, color=WHITE)

    _round_rect(s, Inches(0.45), Inches(1.35), Inches(6.0), Inches(5.3), BG_CARD)
    _textbox(s, Inches(0.7), Inches(1.55), Inches(5.5), Inches(0.4), "Detection pipeline", size=16, bold=True, color=ACCENT_2)
    _multitext(
        s,
        Inches(0.7),
        Inches(2.15),
        Inches(5.5),
        Inches(4.2),
        [
            ("1. Untrusted input → normalization", {"size": 14}),
            ("2. Deterministic rules (9 taxonomy detectors)", {"size": 14}),
            ("3. Groq Prompt Guard + Semantic Safeguard", {"size": 14}),
            ("4. Fusion → Risk (0–100) → Policy", {"size": 14}),
            ("5. ALLOW  |  REVIEW  |  BLOCK", {"size": 14, "bold": True, "color": ACCENT}),
            (" ", {"size": 8}),
            ("Provider failure → UNAVAILABLE / UNCERTAIN", {"size": 13, "color": MUTED}),
            ("Never silent BENIGN", {"size": 14, "bold": True, "color": WARN}),
        ],
        spacing=6,
    )

    _round_rect(s, Inches(6.8), Inches(1.35), Inches(6.0), Inches(5.3), BG_CARD)
    _textbox(s, Inches(7.05), Inches(1.55), Inches(5.5), Inches(0.4), "Protect pipeline", size=16, bold=True, color=ACCENT)
    _multitext(
        s,
        Inches(7.05),
        Inches(2.15),
        Inches(5.5),
        Inches(4.2),
        [
            ("1. Agent security workflow", {"size": 14}),
            ("2. Authn → Authz / tenant isolation", {"size": 14}),
            ("3. Tool Firewall (allowlist, params)", {"size": 14}),
            ("4. Approval binding + replay protection", {"size": 14}),
            ("5. Mock MCP gateway (fingerprint / shadow)", {"size": 14}),
            ("6. Audit → Dashboard (hash only)", {"size": 14}),
            (" ", {"size": 8}),
            ("Key idea: detection is only layer one.", {"size": 13, "color": MUTED}),
            ("Unauthorized tool actions still stop.", {"size": 14, "bold": True, "color": ACCENT}),
        ],
        spacing=6,
    )
    _footer(s, 4, total)

    # 5 Demo
    s = prs.slides.add_slide(blank)
    _fill_slide(s)
    _accent_bar(s)
    _textbox(s, Inches(0.5), Inches(0.3), Inches(10), Inches(0.3), "DEMO HIGHLIGHTS", size=12, bold=True, color=ACCENT)
    _textbox(s, Inches(0.5), Inches(0.65), Inches(12), Inches(0.5), "Six scenarios that prove the boundary.", size=26, bold=True, color=WHITE)
    demos = [
        ("1  Benign request", "ALLOW / REVIEW — safe path", ACCENT),
        ("2  Direct injection", "BLOCK / REVIEW — tools not executed", DANGER),
        ("3  Indirect injection", "Untrusted document content contained", WARN),
        ("4  Intent hijack", "Tool Firewall DENY", DANGER),
        ("5  High-risk delete", "REQUIRES_APPROVAL (bound)", WARN),
        ("6  MCP tamper / shadow", "Integrity DENY — mock MCP only", DANGER),
    ]
    positions = [(0.5, 1.45), (4.55, 1.45), (8.6, 1.45), (0.5, 4.0), (4.55, 4.0), (8.6, 4.0)]
    for (title, body, color), (lx, ty) in zip(demos, positions):
        _round_rect(s, Inches(lx), Inches(ty), Inches(3.8), Inches(2.2), BG_CARD)
        _rect(s, Inches(lx), Inches(ty), Inches(0.12), Inches(2.2), color)
        _textbox(s, Inches(lx + 0.3), Inches(ty + 0.4), Inches(3.3), Inches(0.5), title, size=15, bold=True, color=WHITE)
        _textbox(s, Inches(lx + 0.3), Inches(ty + 1.05), Inches(3.3), Inches(0.8), body, size=14, color=MUTED)
    _footer(s, 5, total)

    # 6 Metrics
    s = prs.slides.add_slide(blank)
    _fill_slide(s)
    _accent_bar(s)
    _textbox(s, Inches(0.5), Inches(0.3), Inches(10), Inches(0.3), "EVIDENCE", size=12, bold=True, color=ACCENT)
    _textbox(s, Inches(0.5), Inches(0.65), Inches(12), Inches(0.5), "Evaluation-dataset measurements — not production accuracy.", size=22, bold=True, color=WHITE)
    metrics = [
        ("0.9706", "Live detection F1", "P 0.9925 · R 0.9496"),
        ("45/45", "Security invariants", "INV-01 … INV-45"),
        ("11/11", "Authorization eval", "Phase 16 matrix"),
        ("15/15", "MCP eval", "Mock gateway only"),
        ("371", "Backend tests", "pytest suite"),
        ("Hash-only", "Audit trail", "No raw prompts / tokens"),
    ]
    positions = [(0.5, 1.45), (4.55, 1.45), (8.6, 1.45), (0.5, 4.05), (4.55, 4.05), (8.6, 4.05)]
    for (big, label, sub), (lx, ty) in zip(metrics, positions):
        _round_rect(s, Inches(lx), Inches(ty), Inches(3.8), Inches(2.15), BG_CARD)
        _textbox(s, Inches(lx + 0.25), Inches(ty + 0.25), Inches(3.3), Inches(0.7), big, size=28, bold=True, color=ACCENT)
        _textbox(s, Inches(lx + 0.25), Inches(ty + 1.0), Inches(3.3), Inches(0.4), label, size=14, bold=True, color=WHITE)
        _textbox(s, Inches(lx + 0.25), Inches(ty + 1.45), Inches(3.3), Inches(0.4), sub, size=12, color=MUTED)
    _footer(s, 6, total)

    # 7 Why
    s = prs.slides.add_slide(blank)
    _fill_slide(s)
    _accent_bar(s)
    _round_rect(s, Inches(8.3), Inches(1.5), Inches(4.5), Inches(4.5), BG_CARD)
    _picture(s, ASSETS / "shield.jpg", Inches(8.45), Inches(1.65), width=Inches(4.2), height=Inches(4.2))
    _textbox(s, Inches(0.5), Inches(0.35), Inches(8), Inches(0.3), "WHY THIS MATTERS", size=12, bold=True, color=ACCENT)
    _textbox(
        s,
        Inches(0.5),
        Inches(0.75),
        Inches(7.5),
        Inches(1.3),
        "Detecting injection text\nis necessary — not sufficient.",
        size=28,
        bold=True,
        color=WHITE,
    )
    points = [
        "Fail-closed on auth, unknown tools, tenant mismatch, replay",
        "Uncertainty → REVIEW, never silent BENIGN",
        "Policy BLOCK cannot be overridden by approval",
        "Tool Firewall is the final execution boundary",
        "MCP / tool outputs remain untrusted",
        "Conservative REVIEW volume is intentional",
    ]
    _multitext(
        s,
        Inches(0.5),
        Inches(2.4),
        Inches(7.5),
        Inches(4.0),
        [(f"●  {p}", {"size": 15, "color": TEXT}) for p in points],
        spacing=8,
    )
    _footer(s, 7, total)

    # 8 Closing
    s = prs.slides.add_slide(blank)
    _fill_slide(s)
    _picture(s, ASSETS / "abstract.jpg", Inches(0), Inches(0), width=SLIDE_W, height=SLIDE_H)
    # Dark overlay panels
    _rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, RGBColor(0x06, 0x0A, 0x14))
    _textbox(
        s,
        Inches(0.8),
        Inches(1.6),
        Inches(11.7),
        Inches(1.2),
        "Protect the tool boundary —\neven when the model is confused.",
        size=30,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    _textbox(
        s,
        Inches(1.8),
        Inches(3.3),
        Inches(9.7),
        Inches(0.9),
        "AegisAI is a hackathon prototype, not a production security product.\nMock MCP and mock tools keep the demo safe while proving the architecture.",
        size=14,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    _round_rect(s, Inches(4.0), Inches(4.5), Inches(5.3), Inches(0.7), BG_CARD)
    _textbox(
        s,
        Inches(4.0),
        Inches(4.62),
        Inches(5.3),
        Inches(0.5),
        "github.com/VarunPunzuri99/AegisAI",
        size=14,
        bold=True,
        color=ACCENT,
        align=PP_ALIGN.CENTER,
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(5.6),
        Inches(11.7),
        Inches(0.4),
        "Detect → Understand → Decide → Protect → Audit",
        size=16,
        bold=True,
        color=ACCENT,
        align=PP_ALIGN.CENTER,
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(6.3),
        Inches(11.7),
        Inches(0.4),
        "ET AI Hackathon 2026 — Agentic Edition  ·  Thank you",
        size=12,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    prs.save(OUT_PPTX)
    return OUT_PPTX


def build_pdf() -> Path:
    """Mirror the deck as a landscape PDF for form upload."""
    W, H = 13.333 * 72, 7.5 * 72  # points (16:9)
    c = canvas.Canvas(str(OUT_PDF), pagesize=(W, H))

    def bg():
        c.setFillColorRGB(*R_BG)
        c.rect(0, 0, W, H, fill=1, stroke=0)

    def accent_bar():
        c.setFillColorRGB(*R_ACCENT)
        c.rect(0, 0, 8, H, fill=1, stroke=0)

    def footer(n, total=8):
        c.setFillColorRGB(*R_MUTED)
        c.setFont("Helvetica", 9)
        c.drawString(36, 22, "AegisAI  ·  ET AI Hackathon 2026 — Agentic Edition  ·  Problem 2")
        c.drawRightString(W - 36, 22, f"{n} / {total}")

    def card(x, y, w, h):
        c.setFillColorRGB(*R_CARD)
        c.roundRect(x, y, w, h, 10, fill=1, stroke=0)

    def draw_image(path: Path, x, y, w, h):
        if path.exists():
            c.drawImage(str(path), x, y, width=w, height=h, preserveAspectRatio=True, mask="auto")

    # Page 1 — Title
    bg()
    draw_image(ASSETS / "circuit.jpg", W * 0.56, 0, W * 0.44, H)
    c.setFillColorRGB(*R_BG)
    c.rect(0, 0, W * 0.60, H, fill=1, stroke=0)
    draw_image(ASSETS / "brand-mark.jpg", 48, H - 140, 100, 100)
    c.setFillColorRGB(*R_TEXT)
    c.setFont("Helvetica-Bold", 48)
    c.drawString(48, H - 230, "AegisAI")
    c.setFillColorRGB(*R_ACCENT)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(48, H - 265, "Agentic Prompt Injection Firewall")
    c.setFillColorRGB(*R_MUTED)
    c.setFont("Helvetica", 13)
    c.drawString(48, H - 310, "A runtime security layer between untrusted content")
    c.drawString(48, H - 330, "and agent tool actions.")
    c.setFillColorRGB(*R_ACCENT)
    c.rect(48, H - 360, 90, 4, fill=1, stroke=0)
    c.setFillColorRGB(*R_MUTED)
    c.setFont("Helvetica", 12)
    c.drawString(48, H - 400, "ET AI Hackathon 2026 — Agentic Edition")
    c.drawString(48, H - 420, "Problem 2  ·  Pitch Deck")
    c.showPage()

    # Page 2 — Problem
    bg()
    accent_bar()
    c.setFillColorRGB(*R_ACCENT)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(40, H - 40, "THE PROBLEM")
    c.setFillColorRGB(*R_TEXT)
    c.setFont("Helvetica-Bold", 30)
    c.drawString(40, H - 85, "Agents read more than chat.")
    c.setFillColorRGB(*R_MUTED)
    c.setFont("Helvetica", 13)
    y = H - 130
    lines = [
        "AI agents ingest documents, websites, emails, and tool outputs.",
        "",
        "Hidden instructions can try to:",
        "→  Override system intent (prompt injection)",
        "→  Hijack tool calls (email, delete, shell)",
        "→  Steal secrets or credentials",
        "→  Poison MCP tool definitions",
        "",
        "Detection alone is not enough — a compromised plan",
        "must still be stopped before execution.",
    ]
    for i, line in enumerate(lines):
        if "Detection alone" in line or "must still" in line:
            c.setFillColorRGB(*R_WARN)
            c.setFont("Helvetica-Bold", 12)
        elif line.startswith("Hidden"):
            c.setFillColorRGB(*R_TEXT)
            c.setFont("Helvetica-Bold", 13)
        else:
            c.setFillColorRGB(*R_TEXT if line.startswith("→") else R_MUTED)
            c.setFont("Helvetica", 12)
        c.drawString(40, y, line)
        y -= 22
    card(W - 340, 80, 300, 360)
    draw_image(ASSETS / "code.jpg", W - 325, 95, 270, 330)
    footer(2)
    c.showPage()

    # Page 3 — Solution
    bg()
    accent_bar()
    c.setFillColorRGB(*R_ACCENT)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(40, H - 40, "THE SOLUTION")
    c.setFillColorRGB(*R_TEXT)
    c.setFont("Helvetica-Bold", 22)
    c.drawString(40, H - 80, "AegisAI sits between untrusted content and agent actions.")
    c.setFillColorRGB(*R_MUTED)
    c.setFont("Helvetica", 12)
    c.drawString(40, H - 110, "Runtime security layer — the frontend never decides ALLOW / BLOCK.")
    stages = [
        ("01", "DETECT", "Rules + Prompt Guard + Safeguard"),
        ("02", "UNDERSTAND", "Evidence fusion & risk scoring"),
        ("03", "DECIDE", "ALLOW · REVIEW · BLOCK"),
        ("04", "PROTECT", "Authz + Tool Firewall + MCP"),
        ("05", "AUDIT", "Hashed events, no raw prompts"),
    ]
    card_w = 170
    gap = 18
    start_x = 40
    for i, (num, title, body) in enumerate(stages):
        x = start_x + i * (card_w + gap)
        card(x, 120, card_w, 280)
        c.setFillColorRGB(*R_ACCENT)
        c.rect(x, 390, card_w, 6, fill=1, stroke=0)
        c.setFont("Helvetica-Bold", 11)
        c.drawString(x + 14, 360, num)
        c.setFillColorRGB(*R_TEXT)
        c.setFont("Helvetica-Bold", 13)
        c.drawString(x + 14, 330, title)
        c.setFillColorRGB(*R_MUTED)
        c.setFont("Helvetica", 10)
        # wrap body
        words = body.split()
        line = ""
        yy = 300
        for w in words:
            test = (line + " " + w).strip()
            if c.stringWidth(test, "Helvetica", 10) > card_w - 28:
                c.drawString(x + 14, yy, line)
                yy -= 14
                line = w
            else:
                line = test
        if line:
            c.drawString(x + 14, yy, line)
    footer(3)
    c.showPage()

    # Page 4 — Architecture
    bg()
    accent_bar()
    c.setFillColorRGB(*R_ACCENT)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(40, H - 40, "ARCHITECTURE")
    c.setFillColorRGB(*R_TEXT)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(40, H - 75, "Defense in depth — detect, then stop the tool boundary.")
    card(36, 70, 420, 380)
    card(W / 2 + 10, 70, 420, 380)
    c.setFillColorRGB(*R_ACCENT2)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(56, 420, "Detection pipeline")
    c.setFillColorRGB(*R_ACCENT)
    c.drawString(W / 2 + 30, 420, "Protect pipeline")
    left = [
        "1. Untrusted input → normalization",
        "2. Deterministic rules (9 detectors)",
        "3. Groq Prompt Guard + Safeguard",
        "4. Fusion → Risk → Policy",
        "5. ALLOW | REVIEW | BLOCK",
        "",
        "Provider failure → UNCERTAIN",
        "Never silent BENIGN",
    ]
    right = [
        "1. Agent security workflow",
        "2. Authn → Authz / tenant",
        "3. Tool Firewall",
        "4. Approval + replay protection",
        "5. Mock MCP gateway",
        "6. Audit → Dashboard",
        "",
        "Unauthorized tool actions still stop.",
    ]
    c.setFont("Helvetica", 12)
    y = 380
    for line in left:
        c.setFillColorRGB(*R_WARN if "Never" in line else (R_ACCENT if "ALLOW" in line else R_TEXT))
        c.drawString(56, y, line)
        y -= 28
    y = 380
    for line in right:
        c.setFillColorRGB(*R_ACCENT if "Unauthorized" in line else R_TEXT)
        c.setFont("Helvetica-Bold" if "Unauthorized" in line else "Helvetica", 12)
        c.drawString(W / 2 + 30, y, line)
        y -= 28
    footer(4)
    c.showPage()

    # Page 5 — Demo
    bg()
    accent_bar()
    c.setFillColorRGB(*R_ACCENT)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(40, H - 40, "DEMO HIGHLIGHTS")
    c.setFillColorRGB(*R_TEXT)
    c.setFont("Helvetica-Bold", 22)
    c.drawString(40, H - 75, "Six scenarios that prove the boundary.")
    demos = [
        ("1  Benign request", "ALLOW / REVIEW — safe path", R_ACCENT),
        ("2  Direct injection", "BLOCK / REVIEW — tools not executed", R_DANGER),
        ("3  Indirect injection", "Untrusted document content contained", R_WARN),
        ("4  Intent hijack", "Tool Firewall DENY", R_DANGER),
        ("5  High-risk delete", "REQUIRES_APPROVAL (bound)", R_WARN),
        ("6  MCP tamper / shadow", "Integrity DENY — mock MCP only", R_DANGER),
    ]
    coords = [
        (40, 300),
        (340, 300),
        (640, 300),
        (40, 80),
        (340, 80),
        (640, 80),
    ]
    for (title, body, color), (x, y) in zip(demos, coords):
        card(x, y, 280, 180)
        c.setFillColorRGB(*color)
        c.rect(x, y, 8, 180, fill=1, stroke=0)
        c.setFillColorRGB(*R_TEXT)
        c.setFont("Helvetica-Bold", 13)
        c.drawString(x + 24, y + 130, title)
        c.setFillColorRGB(*R_MUTED)
        c.setFont("Helvetica", 11)
        c.drawString(x + 24, y + 95, body)
    footer(5)
    c.showPage()

    # Page 6 — Metrics
    bg()
    accent_bar()
    c.setFillColorRGB(*R_ACCENT)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(40, H - 40, "EVIDENCE")
    c.setFillColorRGB(*R_TEXT)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(40, H - 75, "Evaluation-dataset measurements — not production accuracy.")
    metrics = [
        ("0.9706", "Live detection F1", "P 0.9925 · R 0.9496"),
        ("45/45", "Security invariants", "INV-01 … INV-45"),
        ("11/11", "Authorization eval", "Phase 16 matrix"),
        ("15/15", "MCP eval", "Mock gateway only"),
        ("371", "Backend tests", "pytest suite"),
        ("Hash-only", "Audit trail", "No raw prompts / tokens"),
    ]
    coords = [(40, 300), (340, 300), (640, 300), (40, 80), (340, 80), (640, 80)]
    for (big, label, sub), (x, y) in zip(metrics, coords):
        card(x, y, 280, 180)
        c.setFillColorRGB(*R_ACCENT)
        c.setFont("Helvetica-Bold", 26)
        c.drawString(x + 20, y + 115, big)
        c.setFillColorRGB(*R_TEXT)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x + 20, y + 75, label)
        c.setFillColorRGB(*R_MUTED)
        c.setFont("Helvetica", 10)
        c.drawString(x + 20, y + 50, sub)
    footer(6)
    c.showPage()

    # Page 7 — Why
    bg()
    accent_bar()
    c.setFillColorRGB(*R_ACCENT)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(40, H - 40, "WHY THIS MATTERS")
    c.setFillColorRGB(*R_TEXT)
    c.setFont("Helvetica-Bold", 24)
    c.drawString(40, H - 90, "Detecting injection text")
    c.drawString(40, H - 120, "is necessary — not sufficient.")
    points = [
        "Fail-closed on auth, unknown tools, tenant mismatch, replay",
        "Uncertainty → REVIEW, never silent BENIGN",
        "Policy BLOCK cannot be overridden by approval",
        "Tool Firewall is the final execution boundary",
        "MCP / tool outputs remain untrusted",
        "Conservative REVIEW volume is intentional",
    ]
    y = H - 180
    c.setFont("Helvetica", 13)
    for p in points:
        c.setFillColorRGB(*R_TEXT)
        c.drawString(40, y, f"●  {p}")
        y -= 32
    card(W - 340, 90, 300, 300)
    draw_image(ASSETS / "shield.jpg", W - 325, 105, 270, 270)
    footer(7)
    c.showPage()

    # Page 8 — Close
    bg()
    draw_image(ASSETS / "abstract.jpg", 0, 0, W, H)
    c.setFillColorRGB(0.02, 0.04, 0.08)
    c.rect(0, 0, W, H, fill=1, stroke=0)
    c.setFillColorRGB(*R_TEXT)
    c.setFont("Helvetica-Bold", 26)
    c.drawCentredString(W / 2, H / 2 + 80, "Protect the tool boundary —")
    c.drawCentredString(W / 2, H / 2 + 45, "even when the model is confused.")
    c.setFillColorRGB(*R_MUTED)
    c.setFont("Helvetica", 12)
    c.drawCentredString(W / 2, H / 2 - 10, "AegisAI is a hackathon prototype, not a production security product.")
    c.drawCentredString(W / 2, H / 2 - 30, "Mock MCP and mock tools keep the demo safe while proving the architecture.")
    card(W / 2 - 190, H / 2 - 100, 380, 40)
    c.setFillColorRGB(*R_ACCENT)
    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(W / 2, H / 2 - 85, "github.com/VarunPunzuri99/AegisAI")
    c.setFont("Helvetica-Bold", 14)
    c.drawCentredString(W / 2, 80, "Detect → Understand → Decide → Protect → Audit")
    c.setFillColorRGB(*R_MUTED)
    c.setFont("Helvetica", 11)
    c.drawCentredString(W / 2, 50, "ET AI Hackathon 2026 — Agentic Edition  ·  Thank you")
    c.showPage()

    c.save()
    return OUT_PDF


def main() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    pptx_path = build_pptx()
    pdf_path = build_pdf()
    print(f"Wrote {pptx_path}")
    print(f"Wrote {pdf_path}")


if __name__ == "__main__":
    main()
